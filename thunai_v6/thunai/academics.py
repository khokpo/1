"""Courses, internal/semester marks (course staff), student marks view, and HOD analytics."""
import os, io, csv, json, uuid, statistics as st
from datetime import datetime
from flask import Blueprint, render_template, request, redirect, url_for, flash, g, abort, Response, send_file, current_app
import openpyxl
from models import *
from security import roles_required, audit, notify, has_perm
import mentorship as ms, charts
from analytics import student_signal, ami

acad = Blueprint('acad', __name__)
EXAMS = ['Internal 1', 'Internal 2', 'Internal 3 / Model', 'Assignment', 'Semester Exam']
MIN_GROUP = 5  # aggregates (course averages) are shown to students only when at least this many marks exist

def cohort(c): return Student.query.filter_by(programme=c.programme, semester=c.semester).order_by(Student.reg_no).all()

def own_course(cid):
    """Staff may open only courses they teach (server-side IDOR guard); HOD may view all."""
    c = db.session.get(Course, cid) or abort(404)
    if g.user.role == 'staff' and c.staff_id != g.user.id: abort(403)
    return c

def course_stats(c):
    out = []
    for e in EXAMS:
        ms_ = Mark.query.filter_by(course_id=c.id, exam_type=e).all(); p = [m.pct for m in ms_]
        out.append(dict(exam=e, n=len(p), avg=round(st.mean(p), 1) if p else None, med=round(st.median(p), 1) if p else None,
                        sd=round(st.pstdev(p), 1) if len(p) > 1 else None, lo=min(p) if p else None, hi=max(p) if p else None, pcts=p))
    return out

# ---------------- admin: courses ----------------
@acad.route('/admin/courses', methods=['GET', 'POST'])
@roles_required('hod', 'office')
def courses():
    if request.method == 'POST':
        sid = request.form.get('staff') or None
        if sid and (db.session.get(User, int(sid)) or abort(404)).role != 'staff': abort(400)
        c = Course(code=request.form['code'].strip()[:20], name=request.form['name'].strip()[:100], programme=request.form['programme'].strip(),
                   semester=request.form['semester'].strip()[:5], staff_id=int(sid) if sid else None)
        if not c.code or not c.name: flash('Course code and name are required.', 'error')
        else: db.session.add(c); audit('course_create', f'{g.user.name} created {c.code}'); db.session.commit(); flash('Course created.', 'ok')
        return redirect(url_for('acad.courses'))
    cs = Course.query.order_by(Course.programme, Course.semester, Course.code).all()
    return render_template('courses.html', cs=cs, staff=User.query.filter_by(role='staff').order_by(User.name).all(), n={c.id: Mark.query.filter_by(course_id=c.id).count() for c in cs})

@acad.route('/admin/courses/<int:cid>', methods=['POST'])
@roles_required('hod', 'office')
def course_edit(cid):
    c = db.session.get(Course, cid) or abort(404); act = request.form.get('act')
    if act == 'delete':
        if g.user.role != 'hod' or Mark.query.filter_by(course_id=c.id).count(): flash('Only the HOD can delete a course, and only when it has no marks.', 'error')
        else: db.session.delete(c); flash('Course deleted.', 'ok')
    else:
        sid = request.form.get('staff') or None
        if sid and (db.session.get(User, int(sid)) or abort(404)).role != 'staff': abort(400)
        c.name, c.code, c.programme, c.semester = request.form['name'].strip()[:100], request.form['code'].strip()[:20], request.form['programme'].strip(), request.form['semester'].strip()[:5]
        c.staff_id = int(sid) if sid else None; c.active = request.form.get('active') == '1'
        if c.staff_id: notify(c.staff_id, f'You are assigned to {c.code} {c.name}.')
    audit('course_' + str(act), f'{g.user.name} {c.code}'); db.session.commit(); return redirect(url_for('acad.courses'))

# ---------------- course staff ----------------
@acad.route('/acad')
@roles_required('staff')
def dashboard():
    rows = []
    for c in Course.query.filter_by(staff_id=g.user.id).all():
        stats = course_stats(c); rows.append(dict(c=c, n=len(cohort(c)), stats=stats))
    return render_template('course_dashboard.html', rows=rows)

@acad.route('/course/<int:cid>')
@roles_required('staff', 'hod')
def course_view(cid):
    c = own_course(cid); studs = cohort(c); stats = course_stats(c)
    marks = {(m.student_id, m.exam_type): m for m in Mark.query.filter_by(course_id=c.id)}
    avg = charts.bar([s['exam'] for s in stats if s['n']], [('Average %', [s['avg'] for s in stats if s['n']])], 'Average by exam (%)', ymax=100)
    hists = [(s['exam'], charts.hist(s['pcts'], f"{s['exam']} — score distribution (%)")) for s in stats if s['n']]
    return render_template('course_view.html', c=c, studs=studs, stats=stats, marks=marks, EXAMS=EXAMS, avg=avg, hists=hists, can_edit=has_perm(g.user, 'marks_upload', c.id))

@acad.route('/course/<int:cid>/template.xlsx')
@roles_required('staff')
def course_template(cid):
    c = own_course(cid)
    if not has_perm(g.user, 'marks_upload', c.id): abort(403)
    wb = openpyxl.Workbook(); ws = wb.active; ws.append(['reg_no', 'name', 'score'])
    for s in cohort(c): ws.append([s.reg_no, s.name, None])
    b = io.BytesIO(); wb.save(b); b.seek(0); return send_file(b, as_attachment=True, download_name=f'{c.code}_marks_template.xlsx')

def _jobdir():
    d = os.path.join(current_app.instance_path, 'uploads'); os.makedirs(d, exist_ok=True); return d

@acad.route('/course/<int:cid>/upload', methods=['POST'])
@roles_required('staff')
def marks_upload(cid):
    c = own_course(cid); f = request.files.get('file'); exam = request.form.get('exam')
    if not has_perm(g.user, 'marks_upload', c.id, exam): abort(403)   # permission is per course and (optionally) per exam
    try: mx = float(request.form.get('max_score') or 100); assert mx > 0
    except (ValueError, AssertionError): flash('Maximum marks must be a positive number.', 'error'); return redirect(url_for('acad.course_view', cid=cid))
    if exam not in EXAMS or not f or not f.filename.lower().endswith('.xlsx'): flash('Choose an exam type and an .xlsx file.', 'error'); return redirect(url_for('acad.course_view', cid=cid))
    try: rows = list(openpyxl.load_workbook(f, read_only=True, data_only=True).active.iter_rows(values_only=True))
    except Exception: flash('That file could not be read as an Excel workbook.', 'error'); return redirect(url_for('acad.course_view', cid=cid))
    head = [str(h or '').strip().lower() for h in rows[0]] if rows else []
    if 'reg_no' not in head or 'score' not in head: flash('The file needs the columns: reg_no, score.', 'error'); return redirect(url_for('acad.course_view', cid=cid))
    regs = {s.reg_no for s in cohort(c)}; good, bad, seen = [], [], set()
    for i, r in enumerate(rows[1:], start=2):
        d = dict(zip(head, r))
        if not any(v not in (None, '') for v in d.values()): continue
        reg = str(d.get('reg_no') or '').strip().upper(); err = []
        if reg not in regs: err.append('student is not in this course cohort')
        if reg in seen: err.append('duplicate row in file')
        seen.add(reg)
        try:
            sc = float(d.get('score'))
            if not 0 <= sc <= mx: err.append(f'score must be between 0 and {mx:g}')
        except (TypeError, ValueError): err.append('score is blank or not a number (absent students are not imported)'); sc = None
        warn = []
        if not err:
            s_ = Student.query.filter_by(reg_no=reg).first()
            if Mark.query.filter_by(student_id=s_.id, course_id=c.id, exam_type=exam).first(): warn.append('existing mark will be overwritten')
            if sc == 0: warn.append('score is zero — check this is not an absent student')
        (bad if err else good).append(dict(row=i, reg=reg, score=sc, errors='; '.join(err), warn='; '.join(warn)))
    tok = uuid.uuid4().hex
    json.dump(dict(user=g.user.id, course=c.id, exam=exam, max=mx, good=good, bad=bad), open(os.path.join(_jobdir(), tok + '_m.json'), 'w'))
    return render_template('marks_preview.html', c=c, tok=tok, exam=exam, mx=mx, good=good, bad=bad, warns=[x for x in good if x['warn']], received=len(good) + len(bad))

def _job(tok, cid):
    if not tok.isalnum(): abort(404)
    p = os.path.join(_jobdir(), tok + '_m.json')
    if not os.path.exists(p): abort(404)
    j = json.load(open(p))
    if j['user'] != g.user.id or j['course'] != cid: abort(403)
    return j, p

@acad.route('/course/<int:cid>/upload/<tok>/errors.csv')
@roles_required('staff')
def marks_errors(cid, tok):
    own_course(cid); j, _ = _job(tok, cid); out = io.StringIO(); w = csv.writer(out); w.writerow(['Excel row', 'reg_no', 'Issue'])
    for b in j['bad']: w.writerow([b['row'], b['reg'], b['errors']])
    return Response(out.getvalue(), mimetype='text/csv', headers={'Content-Disposition': 'attachment; filename=marks_errors.csv'})

def save_mark(c, reg, exam, score, mx):
    s = Student.query.filter_by(reg_no=reg).first()
    m = Mark.query.filter_by(student_id=s.id, course_id=c.id, exam_type=exam).first() or Mark(student_id=s.id, course_id=c.id, exam_type=exam)
    new = m.id is None; m.score, m.max_score, m.uploaded_by, m.updated = score, mx, g.user.id, datetime.utcnow(); db.session.add(m)
    notify(s.user_id, f'{exam} marks for {c.code} are now available.'); return new

@acad.route('/course/<int:cid>/upload/<tok>/confirm', methods=['POST'])
@roles_required('staff')
def marks_confirm(cid, tok):
    c = own_course(cid); j, p = _job(tok, cid); n = u = 0
    if not has_perm(g.user, 'marks_upload', c.id, j['exam']): abort(403)
    for x in j['good']: n, u = (n + 1, u) if save_mark(c, x['reg'], j['exam'], x['score'], j['max']) else (n, u + 1)
    audit('marks_upload', f'{g.user.name} {c.code} {j["exam"]}: {n} new, {u} updated, {len(j["bad"])} rejected'); db.session.commit(); os.remove(p)
    flash(f'{j["exam"]} imported: {n} new, {u} updated, {len(j["bad"])} rows rejected.', 'ok'); return redirect(url_for('acad.course_view', cid=cid))

@acad.route('/course/<int:cid>/mark', methods=['POST'])
@roles_required('staff')
def mark_edit(cid):
    c = own_course(cid); reg = request.form.get('reg_no', '').upper(); exam = request.form.get('exam')
    if not has_perm(g.user, 'marks_upload', c.id, exam): abort(403)
    try: sc, mx = float(request.form['score']), float(request.form.get('max_score') or 100); assert exam in EXAMS and 0 <= sc <= mx and mx > 0
    except (ValueError, AssertionError, KeyError): flash('Invalid score or maximum.', 'error'); return redirect(url_for('acad.course_view', cid=cid))
    if reg not in {s.reg_no for s in cohort(c)}: abort(403)
    save_mark(c, reg, exam, sc, mx); audit('mark_edit', f'{g.user.name} {c.code} {reg} {exam}'); db.session.commit()
    flash('Mark saved.', 'ok'); return redirect(url_for('acad.course_view', cid=cid))

# ---------------- student: own marks only ----------------
def marks_cfgs(student_id):
    """Marks + chart configs for ONE student. Course averages are aggregates, shown only when the group is large enough."""
    ms_ = Mark.query.filter_by(student_id=student_id).all(); by = {}
    for m in ms_: by.setdefault(m.course, {})[m.exam_type] = m
    courses_ = sorted(by, key=lambda c: c.code); labels = [c.code for c in courses_]; avgs = {}
    for c in courses_:
        for e in EXAMS:
            allm = Mark.query.filter_by(course_id=c.id, exam_type=e).all()
            avgs[(c.id, e)] = round(st.mean(x.pct for x in allm), 1) if len(allm) >= MIN_GROUP else None
    sets = [(e, [by[c][e].pct if e in by[c] else None for c in courses_]) for e in EXAMS if any(e in by[c] for c in courses_)]
    cfg = charts.bar(labels, sets, 'My marks by course (%)', ymax=100) if sets else None
    return by, courses_, avgs, cfg

@acad.route('/student/marks')
@roles_required('student')
def my_marks():
    s = Student.query.filter_by(user_id=g.user.id).first_or_404()
    by, cs, avgs, cfg = marks_cfgs(s.id); sig = student_signal(s.id)
    trend = None
    if sig['recs']:
        L = [f'P{r.period}' for r in sig['recs']]
        trend = charts.line(L, [('Attendance %', [r.attendance for r in sig['recs']]), ('Assessment %', [r.assessment for r in sig['recs']]),
                                ('Assignments %', [r.assignment_pct for r in sig['recs']]), ('Academic Momentum', [ami(r) for r in sig['recs']])], 'My progress over time')
    return render_template('my_marks.html', s=s, by=by, cs=cs, avgs=avgs, cfg=cfg, trend=trend, EXAMS=EXAMS, MIN_GROUP=MIN_GROUP)

# ---------------- HOD analytics ----------------
@acad.route('/admin/analytics')
@roles_required('hod')
def analytics_page():
    studs = Student.query.all(); sigs = [student_signal(s.id) for s in studs]
    lv = ['Stable', 'Monitor', 'Attention Recommended', 'Priority Support']; cnt = [sum(1 for x in sigs if x['level'] == l) for l in lv]
    periods = sorted({r.period for r in AcademicRecord.query}); per = {p: AcademicRecord.query.filter_by(period=p).all() for p in periods}
    mean = lambda v: round(st.mean(v), 1) if v else None
    C = {}
    C['status'] = charts.doughnut(lv, cnt, ['#059669', '#2563eb', '#d97706', '#dc2626'], 'Students by support level')
    C['trend'] = charts.line([f'P{p}' for p in periods], [('Attendance %', [mean([r.attendance for r in per[p]]) for p in periods]), ('Assessment %', [mean([r.assessment for r in per[p]]) for p in periods]),
                                                           ('Assignments %', [mean([r.assignment_pct for r in per[p]]) for p in periods])], 'Department averages by period')
    C['ami'] = charts.hist([x['ami'] for x in sigs if x['ami'] is not None], 'Academic Momentum distribution')
    cs = Course.query.order_by(Course.code).all(); lab = [c.code for c in cs]
    ex = [(e, [mean([m.pct for m in Mark.query.filter_by(course_id=c.id, exam_type=e)]) for c in cs]) for e in EXAMS]
    C['courses'] = charts.bar(lab, [x for x in ex if any(v is not None for v in x[1])], 'Course averages by exam (%)', ymax=100) if cs else None
    cats = {}
    for r in SupportRequest.query: cats[r.category] = cats.get(r.category, 0) + 1
    C['req'] = charts.bar(list(cats), [('Requests', list(cats.values()))], 'Support requests by category', horizontal=True) if cats else None
    its = {}
    for i in Intervention.query: its[i.itype] = its.get(i.itype, 0) + 1
    C['iv'] = charts.bar(list(its), [('Interventions', list(its.values()))], 'Interventions by type', horizontal=True) if its else None
    mentors = User.query.filter_by(role='mentor').all()
    C['load'] = charts.bar([m.name for m in mentors], [('Students', [len(ms.students_of(m.id)) for m in mentors])], 'Students per mentor (allocation view)') if mentors else None
    return render_template('analytics.html', C=C, n=len(studs))
